"""Render synthetic knowledge sources into the supported file formats."""

from __future__ import annotations

import json
import os
import time
from collections.abc import Callable
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from shopmate_ecommerce_assistant.core.ids import stable_uuid
from shopmate_ecommerce_assistant.synthetic.catalog_data import Dataset
from shopmate_ecommerce_assistant.synthetic.eval_cases import build_eval_cases
from shopmate_ecommerce_assistant.synthetic.knowledge_corpus import (
    PolicyDoc,
    dumps_rows,
    faq_records,
    manifest_payload,
)

_FONT_CANDIDATES = (
    os.environ.get("SHOPMATE_PDF_FONT", ""),
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\simsun.ttc",
    "/System/Library/Fonts/PingFang.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
)

_font_name: str | None = None


def cjk_font() -> str:
    """Return a registered font name that can render Chinese.

    A TrueType font is preferred because it embeds a Unicode map that keeps PDF text
    extractable. When no system font is available the Adobe CID font is used instead, which
    renders correctly but may not extract cleanly.
    """
    global _font_name
    if _font_name is not None:
        return _font_name
    for candidate in _FONT_CANDIDATES:
        if not candidate:
            continue
        path = Path(candidate)
        if not path.is_file():
            continue
        try:
            pdfmetrics.registerFont(TTFont("ShopMateCJK", str(path)))
        except Exception:
            continue
        _font_name = "ShopMateCJK"
        return _font_name
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _font_name = "STSong-Light"
    return _font_name


def _markdown(doc: PolicyDoc) -> str:
    lines = [f"# {doc.title}", ""]
    for section in doc.sections:
        lines.extend([f"## {section.heading}", ""])
        for paragraph in section.paragraphs:
            lines.extend([paragraph, ""])
    return "\n".join(lines).rstrip() + "\n"


def _plain_text(doc: PolicyDoc) -> str:
    lines = [doc.title, ""]
    for section in doc.sections:
        lines.extend([f"【{section.heading}】", ""])
        for paragraph in section.paragraphs:
            lines.extend([paragraph, ""])
    return "\n".join(lines).rstrip() + "\n"


def _pdf(doc: PolicyDoc, target: Path) -> None:
    font = cjk_font()
    page_width, page_height = A4
    left = 20 * mm
    right = page_width - 20 * mm
    top = page_height - 22 * mm
    bottom = 25 * mm

    pdf = canvas.Canvas(str(target), pagesize=A4, invariant=1)
    pdf.setTitle(doc.title)
    y = top

    def new_page() -> float:
        pdf.showPage()
        return top

    def write(text: str, size: float, leading: float, *, gap: float = 2 * mm) -> None:
        nonlocal y
        for line in simpleSplit(text, font, size, right - left):
            if y < bottom:
                y = new_page()
            pdf.setFont(font, size)
            pdf.drawString(left, y, line)
            y -= leading
        y -= gap

    write(doc.title, 16, 22)
    write(f"版本 {doc.version} · 生效日期 {doc.effective_from[:10]}", 9, 14, gap=6 * mm)
    for section in doc.sections:
        if section.heading.startswith("附录"):
            y = new_page()
        elif y < bottom + 30 * mm:
            y = new_page()
        write(section.heading, 12.5, 18)
        for paragraph in section.paragraphs:
            write(paragraph, 10.5, 16)
        y -= 3 * mm
    pdf.save()


def _docx(doc: PolicyDoc, target: Path) -> None:
    document = DocxDocument()
    document.add_heading(doc.title, level=0)
    meta = document.add_paragraph(f"版本 {doc.version} · 生效日期 {doc.effective_from[:10]}")
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_added = False
    for section in doc.sections:
        document.add_heading(section.heading, level=2)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        if doc.key == "custom-product-policy" and not table_added:
            table = document.add_table(rows=1, cols=2)
            header = table.rows[0].cells
            header[0].text = "情形"
            header[1].text = "处理方式"
            for case, handling in (
                ("尺寸与确认稿不符", "免费重做或退货"),
                ("颜色存在明显偏差", "按确认稿重做"),
                ("刻字内容错误（非消费者原因）", "免费重做"),
            ):
                row = table.add_row().cells
                row[0].text = case
                row[1].text = handling
            table_added = True
    document.save(str(target))


def _pptx(doc: PolicyDoc, target: Path) -> None:
    presentation = Presentation()
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = doc.title
    cover.placeholders[1].text = f"版本 {doc.version} · 生效日期 {doc.effective_from[:10]}"

    for section in doc.sections:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = section.heading
        body = slide.placeholders[1].text_frame
        body.text = section.paragraphs[0]
        body.paragraphs[0].font.size = Pt(16)
        for paragraph in section.paragraphs[1:]:
            bullet = body.add_paragraph()
            bullet.text = paragraph
            bullet.font.size = Pt(16)
    presentation.save(str(target))


def _retry_write(action: Callable[[], None], *, attempts: int = 4) -> None:
    """Run a binary document write, retrying transient Windows sharing violations.

    Windows antivirus/indexers briefly lock freshly written ``.docx``/``.pdf``/``.pptx``
    files, which surfaces as ``OSError: [Errno 22] Invalid argument`` from ``zipfile``.
    The content is deterministic, so retrying the same write is safe.
    """
    last: OSError | None = None
    for attempt in range(attempts):
        try:
            action()
            return
        except OSError as exc:  # pragma: no cover - depends on the host filesystem
            last = exc
            time.sleep(0.05 * (attempt + 1))
    assert last is not None
    raise last


def render_policy(doc: PolicyDoc, directory: Path) -> Path:
    """Write one policy document in its declared format and return the path."""
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / doc.filename
    if doc.file_format == "md":
        target.write_text(_markdown(doc), encoding="utf-8")
    elif doc.file_format == "txt":
        target.write_text(_plain_text(doc), encoding="utf-8")
    elif doc.file_format == "pdf":
        _retry_write(lambda: _pdf(doc, target))
    elif doc.file_format == "docx":
        _retry_write(lambda: _docx(doc, target))
    elif doc.file_format == "pptx":
        _retry_write(lambda: _pptx(doc, target))
    else:  # pragma: no cover - guarded by the corpus definition
        raise ValueError(f"unsupported file format: {doc.file_format}")
    return target


def render_all(docs: tuple[PolicyDoc, ...], knowledge_dir: Path) -> list[Path]:
    """Render every policy document under ``knowledge_dir/policy``."""
    policy_dir = knowledge_dir / "policy"
    return [render_policy(doc, policy_dir) for doc in docs]


def write_eval_cases(target: Path) -> int:
    """Write the 100 generated evaluation cases as JSON Lines, unless a dataset exists.

    Seeding must not clobber the curated dataset. Since stage 13 the file on disk is
    hand-extended to 136 cases (authorization / multi-turn / short-query / no-answer
    variants); re-rendering the 100 generated rows here silently deleted those 36 cases
    and made the evaluation report irreproducible. An existing non-empty file is
    therefore left untouched and its real case count is returned.
    """
    cases = build_eval_cases()
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        existing = sum(
            1 for line in target.read_text(encoding="utf-8").splitlines() if line.strip()
        )
        if existing:
            return existing
    target.write_text(dumps_rows(cases), encoding="utf-8")
    return len(cases)


def write_faq(target: Path) -> int:
    """Write the 50 FAQ rows as JSON Lines."""
    rows = faq_records()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(dumps_rows(rows), encoding="utf-8")
    return len(rows)


def write_manifest(target: Path) -> int:
    """Write the knowledge ingest manifest."""
    payload = manifest_payload()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return len(payload["documents"])


def write_catalog_snapshot(dataset: Dataset, target: Path) -> None:
    """Write a small human-readable snapshot of seeded catalog rows.

    The snapshot is derived output; PostgreSQL remains the source of truth.
    """
    lines = [
        "# 合成商品快照（由 scripts/seed.py 生成，PostgreSQL 为准）",
        "",
        "| SKU | 商品 | 类目 | 价格(分) | 库存 |",
        "|---|---|---|---|---|",
    ]
    stock = {row["variant_id"]: row["available_qty"] for row in dataset.inventory}
    for variant in dataset.variants:
        lines.append(
            "| {sku} | {title} | {category} | {price} | {qty} |".format(
                sku=variant["sku"],
                title=variant["title"],
                category=variant["product_category"],
                price=variant["price_cents"],
                qty=stock.get(variant["id"], 0),
            )
        )
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def dataset_fingerprint(dataset: Dataset) -> str:
    """Return a stable fingerprint of the seeded row ids for duplicate detection."""
    ids = [
        str(dataset.shop["id"]),
        *sorted(str(row["id"]) for row in dataset.products),
        *sorted(str(row["id"]) for row in dataset.orders),
    ]
    return str(stable_uuid("dataset", *ids))
