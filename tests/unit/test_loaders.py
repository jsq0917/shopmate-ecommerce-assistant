"""Loader tests: heading paths, page numbers, table text and error mapping."""

from __future__ import annotations

from pathlib import Path

import pytest

from shopmate_ecommerce_assistant.rag.errors import (
    DocumentParseFailed,
    EmptyDocument,
    UnsupportedDocumentType,
)
from shopmate_ecommerce_assistant.rag.ingest.loaders import load_document
from shopmate_ecommerce_assistant.synthetic.knowledge_corpus import PolicyDoc, Section
from shopmate_ecommerce_assistant.synthetic.writers import render_policy


def _policy_doc(*, file_format: str, sections: tuple[Section, ...]) -> PolicyDoc:
    return PolicyDoc(
        key="loader-fixture",
        title="加载器测试政策",
        doc_type="policy",
        region="CN",
        knowledge_base="policy-cn",
        version="2026.10.1",
        effective_from="2026-10-01T00:00:00Z",
        effective_to=None,
        file_format=file_format,
        sections=sections,
    )


SECTIONS = (
    Section("适用范围", ("本政策适用于测试场景，所有数据均为合成数据。",)),
    Section("退货条件", ("商品需保持完好，配件齐全，不影响二次销售。",)),
    Section("退款时效", ("退款在验收通过后一至七个工作日内原路退回。",)),
)


def test_markdown_keeps_nested_heading_paths(tmp_path: Path) -> None:
    source = tmp_path / "policy.md"
    source.write_text(
        "# 售后总则\n\n总则正文内容。\n\n## 退货\n\n退货正文。\n\n### 生鲜类\n\n生鲜正文。\n",
        encoding="utf-8",
    )

    loaded = load_document(source)

    assert [(block.heading_path, block.text) for block in loaded.blocks] == [
        (("售后总则",), "总则正文内容。"),
        (("售后总则", "退货"), "退货正文。"),
        (("售后总则", "退货", "生鲜类"), "生鲜正文。"),
    ]
    assert loaded.page_count == 0


def test_txt_detects_title_and_bracketed_headings(tmp_path: Path) -> None:
    source = tmp_path / "policy.txt"
    source.write_text(
        "包邮门槛说明\n\n【包邮门槛】\n满 99 元包邮。\n\n【不参与包邮】\n大件商品不包邮。\n",
        encoding="utf-8",
    )

    loaded = load_document(source)

    assert loaded.blocks[0].heading_path == ("包邮门槛说明", "包邮门槛")
    assert loaded.blocks[0].text == "满 99 元包邮。"
    assert loaded.blocks[1].heading_path == ("包邮门槛说明", "不参与包邮")
    assert all(block.heading_path[0] == "包邮门槛说明" for block in loaded.blocks)


def test_docx_keeps_headings_and_table_rows(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "custom.docx"
    document = Document()
    document.add_heading("定制商品售后规则", level=0)
    document.add_heading("加工误差", level=2)
    document.add_paragraph("尺寸偏差超出约定公差时按质量问题处理。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "情形"
    table.rows[0].cells[1].text = "处理方式"
    row = table.add_row().cells
    row[0].text = "尺寸不符"
    row[1].text = "免费重做"
    document.save(str(path))

    loaded = load_document(path)

    headings = [block.heading_path for block in loaded.blocks]
    assert ("定制商品售后规则", "加工误差") in headings
    assert any("尺寸不符 | 免费重做" in block.text for block in loaded.blocks)


def test_pptx_uses_slide_numbers_and_cover_title(tmp_path: Path) -> None:
    doc = _policy_doc(file_format="pptx", sections=SECTIONS)
    path = render_policy(doc, tmp_path)

    loaded = load_document(path)

    assert loaded.page_count == len(SECTIONS) + 1
    slide_blocks = [block for block in loaded.blocks if block.page_no and block.page_no > 1]
    assert slide_blocks, "expected content blocks on section slides"
    assert slide_blocks[0].page_no == 2
    assert slide_blocks[0].heading_path == ("加载器测试政策", "适用范围")


def test_pdf_preserves_page_numbers(tmp_path: Path) -> None:
    doc = _policy_doc(
        file_format="pdf",
        sections=SECTIONS
        + (
            Section(
                "附录：常见问题",
                tuple(f"问题{i}：这是第{i}条测试问答内容。" for i in range(1, 12)),
            ),
        ),
    )
    path = render_policy(doc, tmp_path)

    loaded = load_document(path)

    if not any(block.text.strip() for block in loaded.blocks):
        pytest.skip("当前系统缺少可提取文本的中文字体，跳过 PDF 文本断言")
    assert loaded.page_count >= 2
    assert 2 in {block.page_no for block in loaded.blocks}
    assert any(
        block.heading_path and block.heading_path[0] == "加载器测试政策" for block in loaded.blocks
    )


def test_unsupported_suffix_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "policy.csv"
    source.write_text("a,b\n", encoding="utf-8")

    with pytest.raises(UnsupportedDocumentType):
        load_document(source)


def test_empty_markdown_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "empty.md"
    source.write_text("\n\n   \n", encoding="utf-8")

    with pytest.raises(EmptyDocument):
        load_document(source)


def test_corrupt_pdf_is_reported_as_parse_failure(tmp_path: Path) -> None:
    source = tmp_path / "broken.pdf"
    source.write_bytes(b"not a real pdf payload")

    with pytest.raises(DocumentParseFailed):
        load_document(source)


# --- images (stage 20) --------------------------------------------------------------


def _render_screenshot(tmp_path: Path) -> Path:
    """Render the fixture policy as a PNG, skipping when the host has no CJK font."""
    doc = _policy_doc(file_format="png", sections=SECTIONS)
    try:
        return render_policy(doc, tmp_path)
    except RuntimeError as exc:  # pragma: no cover - depends on installed fonts
        pytest.skip(f"当前系统缺少中文字体，跳过图片渲染：{exc}")


def test_image_is_read_through_ocr(tmp_path: Path) -> None:
    """A screenshot yields ordinary text blocks plus its own file metadata."""
    source = _render_screenshot(tmp_path)

    loaded = load_document(source)

    text = "\n".join(block.text for block in loaded.blocks)
    # OCR keeps the wording; assert on a digit-free fragment because engines drop the
    # spaces around numbers ("第 7 日" → "第7日") depending on the model version.
    assert "商品需保持完好" in text
    assert loaded.mime_type == "image/png"
    assert loaded.page_count == 1
    image = loaded.metadata["image"]
    assert image["filename"] == source.name
    assert image["width"] == 1080
    assert image["height"] > 0
    assert image["ocr_mean_confidence"] > 0.5
    assert image["ocr_lines"] >= 3


def test_image_block_uses_page_one_and_no_fake_heading(tmp_path: Path) -> None:
    """A screenshot has no sections; the loader must not invent a heading path."""
    source = _render_screenshot(tmp_path)

    loaded = load_document(source)

    assert {block.page_no for block in loaded.blocks} == {1}
    assert all(block.heading_path == () for block in loaded.blocks)


def test_image_without_legible_text_is_rejected(tmp_path: Path) -> None:
    from PIL import Image

    source = tmp_path / "blank.png"
    Image.new("RGB", (600, 400), "white").save(source)

    with pytest.raises(EmptyDocument):
        load_document(source)


def test_non_image_with_image_suffix_is_rejected(tmp_path: Path) -> None:
    """Magic bytes, not the extension, decide: a renamed text file must not reach OCR."""
    source = tmp_path / "fake.png"
    source.write_bytes(b"this is not a png payload")

    with pytest.raises(DocumentParseFailed):
        load_document(source)


def test_image_above_pixel_limit_is_rejected(tmp_path: Path) -> None:
    from PIL import Image

    from shopmate_ecommerce_assistant.rag.ingest import loaders

    # 9000 × 5000 = 45 M pixels > the 40 M cap; mode "1" keeps the fixture small in memory.
    source = tmp_path / "huge.png"
    Image.new("1", (9000, 5000), 1).save(source)
    assert source.stat().st_size < loaders.IMAGE_MAX_BYTES

    with pytest.raises(DocumentParseFailed):
        load_document(source)


def test_gif_is_rejected_as_unsupported(tmp_path: Path) -> None:
    from PIL import Image

    source = tmp_path / "animation.gif"
    Image.new("RGB", (32, 32), "white").save(source, format="GIF")

    with pytest.raises(UnsupportedDocumentType):
        load_document(source)
